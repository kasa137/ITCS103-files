import tkinter as tk
from tkinter import ttk
import openpyxl as op
from tkinter import messagebox


window = tk.Tk()
window.title("Simple Ordering System")
window.configure(bg="lightblue")

def display():
    workbook = tk.load_workbook("orderDB.xlsx")
    sheet = workbook.active
    
    for content in table.get_children():
        table.delete(content)

    for rows in sheet.iter_rows(main_row = 2,values_only = True):
        table.insert("",tk.END,values=rows)

def validation():
    cn = cname_entry.get()
    p = product_entry.get()
    q = qty_entry.get()
    pc = price_entry.get()

    if not cn or not p or not q or not pc:
        messagebox.showerror("Error","all fiels should be fiilled")

    if not q or not pc.isdigit():
        messagebox.showerror("Error","Quantity and Price should be a number")

def auto_populates(event):
    selected = table.focus()
    values = table.item(selected,"values")

    if values:
        cname_entry.delete(0,tk.END)
        product_entry.delete(0,tk.END)
        qty_entry.delete(0,tk.END)
        price_entry.delete(0,tk.END)

        cname_entry.insert(0,values[1])
        product_entry.insert(0,values[2])
        qty_entry.insert(0,values[3])
        price_entry.insert(0,values[4])

def update():
    selected = table.focus()

    if not selected :
        messagebox.showerror("insert product first ")

    if not validation:
        return
    cn = cname_entry.get()
    p = product_entry.get()
    q = int(qty_entry.get())
    pc = int(price_entry.get())

    total = q * pc 

    values = table.item(selected, "values")
    order_id = values[0]

    workbook = tk.load_workbook("orderDB.xlsx")
    sheet = workbook.active



# Form Title
title = tk.Label(window, text="Simple Ordering System", font=("Times New Roman", 14, "bold"), bg="lightblue")
title.grid(row=0, column=0, columnspan=6)

# Frame
genframe = tk.Frame(window, bg="lightblue", bd=2, relief="groove")
genframe.grid(row=1, column=0, columnspan=7, padx=10, pady=10)

# Customer Name Entry
cname_entry = tk.Entry(genframe, font=("Poppins", 12))
cname_entry.grid(row=2, column=1, columnspan=2, padx=10, pady=(10, 0))

cname_label = tk.Label(genframe, text="Customer Name", font=("Poppins", 10, "italic"), bg="lightblue")
cname_label.grid(row=3, column=1, columnspan=2)

# Product Entry
product_entry = tk.Entry(genframe, font=("Poppins", 12))
product_entry.grid(row=2, column=3, columnspan=2, padx=10, pady=(10, 0))

product_label = tk.Label(genframe, text="Product", font=("Poppins", 10, "italic"), bg="lightblue")
product_label.grid(row=3, column=3, columnspan=2)

# Quantity Entry
qty_entry = tk.Entry(genframe, font=("Poppins", 12))
qty_entry.grid(row=4, column=1, columnspan=2, padx=10, pady=(10, 0))

qty_label = tk.Label(genframe, text="Quantity", font=("Poppins", 10, "italic"), bg="lightblue")
qty_label.grid(row=5, column=1, columnspan=2)

# Price Entry
price_entry = tk.Entry(genframe, font=("Poppins", 12))
price_entry.grid(row=4, column=3, columnspan=2, padx=10, pady=(10, 0))

price_label = tk.Label(genframe, text="Price", font=("Poppins", 10, "italic"), bg="lightblue")
price_label.grid(row=5, column=3, columnspan=2)

# Buttons
submit_btn = tk.Button(window, text="Submit", font=("Poppins", 12, "bold"), bg="lightpink",command = validation)
submit_btn.grid(row=6, column=1, pady=(10, 20))

update_btn = tk.Button(window, text="Update",font=("Poppins", 12, "bold"), bg="lightgreen",command = update)
update_btn.grid(row=6, column=2)

delete_btn = tk.Button(window, text="Delete", bg="red", fg="white",font=("Poppins", 12, "bold"),command = auto_populates)
delete_btn.grid(row=6, column=3)

# Table
table = ttk.Treeview(
    window,
    columns=("Order ID", "Customer Name", "Product", "Quantity", "Price", "Total"),
    show="headings"
)

for headings in ("Order ID", "Customer Name", "Product", "Quantity", "Price", "Total"):
    table.heading(headings, text=headings)

table.grid(row=7, column=0, columnspan=6, padx=10, pady=10)


window.mainloop()

