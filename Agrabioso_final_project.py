#final project
import tkinter as unak
import openpyxl as ts
from tkinter import messagebox,ttk

window = unak.Tk()
window.configure(bg="lightyellow")
window.title("Laboratory Equipment Borrowing System")

def display():
    workbook = ts.load_workbook('Agrabioso_database.xlsx')
    sheet = workbook.active
    
    for ihan in tree.get_children():
        tree.delete(ihan)
        
    for ihan in sheet.iter_rows(min_row=2,values_only=True):
        tree.insert("",unak.END,values=ihan)
        
def validation():
    
    name = ename.get()
    equipment = equip.get()
    quan = equan.get()
    date = edate.get()
    con = condition.get()

    if not name or not date or not quan or not equipment or not con:
        messagebox.showerror("Error", "Please fill in all fields.\n     Ó⁠╭⁠╮⁠Ò")
        return False

    if not quan.isdigit():
        messagebox.showerror("Error", "Quantity must be numbers only.\n    Ó⁠╭⁠╮⁠Ò")
        return False

    return True
    
def create():
    if not validation():
        return
        
    name = ename.get()
    equipment = equip.get()
    quan = equan.get()
    date = edate.get()
    con = condition.get()
    
    workbook = ts.load_workbook('Agrabioso_database.xlsx')
    sheet = workbook.active
    
    id = sheet.max_row
    
    sheet.append([id,name,equipment,quan,date,con])
    workbook.save("Agrabioso_database.xlsx")
    
    messagebox.showinfo("Success","Borrowed Equipment Recorded Succesfully\n     ✧⁠◝⁠(⁠⁰⁠▿⁠⁰⁠)⁠◜⁠✧ ")
    display()
    
def select(event):
    select = tree.focus()
    values = tree.item(select,"values")
    
    if values:
        ename.delete(0,unak.END)
        equip.delete(0,unak.END)
        equan.delete(0,unak.END)
        edate.delete(0,unak.END)
        
        ename.insert(0,values[1])
        equip.insert(0,values[2])
        equan.insert(0,values[3])
        edate.insert(0,values[4])
        condition.set(values[5])
        
def update():
    select = tree.focus()
    
    if not select:
        messagebox.showerror("Error","Select a record first\n      Ó⁠╭⁠╮⁠Ò ")
        return
        
    if not validation():
        return
        
    values = tree.item(select,"values")
    id = int(values[0])
    
    name = ename.get()
    equipment = equip.get()
    quan = equan.get()
    date = edate.get()
    condi = condition.get()
    
    workbook = ts.load_workbook("Agrabioso_database.xlsx")
    sheet = workbook.active
    
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == id :
            row[1].value = name
            row[2].value = equipment
            row[3].value = quan
            row[4].value = date
            row[5].value = condi
            
    workbook.save("Agrabioso_database.xlsx")
    
    messagebox.showinfo("Succes","Borrowed equipment updated succesfully\n     ✧⁠◝⁠(⁠⁰⁠▿⁠⁰⁠)⁠◜⁠✧ ")
    
    display()
    
def delete():
    select = tree.focus()
    
    if not select:
        messagebox.showerror("Error","Select a record first\n    Ó⁠╭⁠╮⁠Ò")
        return
        
    values = tree.item(select,"values")
    id = int(values[0])
    
    sureness = messagebox.askyesno("U sure?","Are you sure you want to delete this record?\n     ༼⁠⁰⁠o⁠⁰⁠；⁠༽")
    if not sureness:
        return
        
    workbook = ts.load_workbook("Agrabioso_database.xlsx")
    sheet = workbook.active
    
    for ihan, row in enumerate(sheet.iter_rows(min_row=2),start=2):
        if row[0].value == id:
            sheet.delete_rows(ihan)
            break
            
    workbook.save("Agrabioso_database.xlsx")
    
    messagebox.showinfo("succes","Record deleted succesfully\n    ✧⁠◝⁠(⁠⁰⁠▿⁠⁰⁠)⁠◜⁠✧")
    display()
    

##########################################################################################################################################

title = unak.Label(window,text="Laboratory Equipment Borrowing System",bg ="lightyellow",font=("Lobster",15))
title.grid(column = 0,row = 0, columnspan = 8)

ename = unak.Entry(window)
ename.grid(column = 3,row = 1,columnspan= 2,padx = 10,pady=10)

lname = unak.Label(window,text="Name",bg="lightyellow",font=("Georgia",10))
lname.grid(column=0,row=1,columnspan=2,padx=10,pady=10)

equip = unak.Entry(window,)
equip.grid(column=3,row=2,columnspan =2,padx=5,pady=10)

lequip = unak.Label(window,text="Equipment",bg="lightyellow",font=("Georgia",10))
lequip.grid(column=0,row=2,columnspan=2,padx=10,pady=10)

equan = unak.Entry(window)
equan.grid(column=3,row=3,columnspan=2,padx=5,pady=10)

quan = unak.Label(window,text="Quantity",bg="lightyellow",font=("Georgia",10))
quan.grid(column=0,row=3,columnspan=2,padx=10,pady=10)

edate = unak.Entry(window)
edate.grid(column=3,row=4,columnspan=2,padx=5,pady=10)

date = unak.Label(window,text="Date",bg="lightyellow",font=("Georgia",10))
date.grid(column=0,row=4,columnspan=2,padx=10,pady=10)

condi = unak.Label(window,text ="Euipment Condition",bg="lightyellow")
condi.grid(column=0,row=5,columnspan=4,padx=10,pady=10)

con = unak.IntVar()
condition = unak.StringVar()

good = unak.Radiobutton(window,text="Good Condition",value="Good",bg="lightyellow",activebackground="lightyellow",variable=condition)
good.grid(column=0,row=6,columnspan=2, padx=10,pady=10)

bad = unak.Radiobutton(window,text="Damaged",value="Damaged",bg="lightyellow",activebackground="lightyellow",variable=condition)
bad.grid(column=3,row=6,columnspan=2,padx=10,pady=10)

update = unak.Button(window,text="Update",relief = "sunken",command=update,activebackground="yellow",activeforeground="black",bg="orange",fg="black")
update.grid(column=0,row=9,columnspan=3,padx=10,pady=10)

dilit = unak.Button(window,text="Delete",relief ="sunken",command=delete, activebackground="violet", activeforeground="black",bg="red",fg="white")
dilit.grid(column=6,row=9,columnspan=3,padx=10,pady=10)

submit = unak.Button(window,text="Borrow",relief ="sunken",command=create, activebackground="lightgreen", activeforeground="black",bg="green",fg="white")
submit.grid(column=4,row=9,columnspan=2,padx=10,pady=10)

treefra=unak.Frame(window,bg = "lightgrey")
treefra.grid(column=5,row=1,columnspan=4,rowspan=4,padx=20,pady=20)

tree = ttk.Treeview(treefra,columns=("ID","NAME","EQUIPMENT","QUANTITY","DATE","CONDITION"),show="headings")

for col in ("ID","NAME","EQUIPMENT","QUANTITY","DATE","CONDITION"):
    tree.heading(col, text=col)
    
tree.grid(column = 5,row = 1,columnspan = 4,rowspan=4,padx=20,pady=20)    
    
tree.bind("<<TreeviewSelect>>",select)

display()

window.mainloop()
